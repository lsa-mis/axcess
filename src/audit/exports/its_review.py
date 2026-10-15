"""The U-M ITS Accessibility Team's "Technical Review" workbook, filled in.

The University of Michigan ITS Accessibility Team records each evaluation in
one spreadsheet template (``templates/its_technical_review.xlsx``, kept
byte-for-byte as the team supplied it). This module opens that template and
fills it from one completed scan, so the team receives its own format rather
than a second workbook to transcribe.

The template stays exactly as the team supplied it: its formulas, dropdowns,
named ranges, styles, example values and reference tables are not changed,
even where they have known faults. Only data cells are written. This is the
team's decision for now (October 2026). Two consequences follow, and the
tests pin both:

* The Summary formulas count rows 2-49 only, so they miss any row past the
  48th. Their ``DIVIDE()`` is a Google Sheets function, so Excel shows
  ``#NAME?`` for the percentages.
* Rows past the template's own (Findings Issues row 49, Test Plan row 138)
  are written as plain cells, with no dropdown or styling. Every occurrence
  is still listed rather than dropped.

What each sheet receives:

* **Coversheet**: the scan date, the site's host as the IT name, its start
  URL, the tools the scan used, and, when an expert evaluation records it,
  the purpose as Background. Tester(s) is left for the team, as the team
  asked: Axcess has no user accounts to name a tester from. Fields only a
  person knows (U-M contact, vendor contact, VPAT, description, login info)
  keep the template's own prompt so they read as unfilled, not as blank
  facts, and IT provenance and IT type keep the template's values.
* **Test Plan**: every page the scan checked, one row each, with its Test
  name and URL; the Description column is left for the team. Column A feeds
  the Test dropdown on Findings Issues through the template's
  ``GenericTestList`` name.
* **Summary**: a short factual narrative in the Narrative Summary cell.
* **Findings Issues**: one row per occurrence (one element on one page) of
  each Axcess issue a person may need to act on, as the ITS team asked: the
  team tracks, fixes and verifies each barrier on its own row. An issue's
  rows sit together, in the app's priority order, and each row's Test is
  the page it is on. Category, Criterion, Impact and "Most affects" use the
  template's own Reference and Reference Impact wording wherever the
  criterion is listed there, so the team's dropdowns and Summary counts
  match. Issues Axcess marks "For information", and occurrences someone
  marked "Not a problem", are left out: the template asks for barriers only.
* **Reference**, **Reference Impact**: unchanged.

Severity is a first estimate. The template's scale judges access to core
and secondary tasks, which a crawl cannot know, so the automated impact
rating is mapped onto it conservatively and nothing is ever marked Blocker.
The Summary narrative says so.

Every value written from the scan (page titles, URLs, element locators) is
stored as text, never as a formula: a page titled ``=HYPERLINK(...)`` must
not become a live formula in the team's spreadsheet.
"""

from __future__ import annotations

import io
import re
import sqlite3
from dataclasses import dataclass
from datetime import datetime
from importlib import resources
from typing import Any
from urllib.parse import urlsplit

from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from audit import coverage_matrix, evaluation
from audit.exports.audit_report import IssueLocation, issue_locations, plain_text
from audit.exports.collector import ExportScan
from audit.labels import REVIEW_LANE_LABELS
from audit.web import issues

_TEMPLATE_PACKAGE = "audit.exports"
_TEMPLATE_PATH = ("templates", "its_technical_review.xlsx")

# The template's sheet names, as the ITS team wrote them.
_COVER = "Coversheet"
_PLAN = "Test Plan"
_SUMMARY = "Summary"
_FINDINGS = "Findings Issues"
_REFERENCE = "Reference"
_IMPACT = "Reference Impact"

# Findings Issues: the template ships 48 numbered blank rows (2-49), and its
# Summary formulas and dropdowns cover only those. The export keeps at least
# that many, so the team can still add issues by hand.
_FIRST_ISSUE_ROW = 2
_TEMPLATE_LAST_ISSUE_ROW = 49
# Test Plan: rows 5-138 carry the template's styles and its named range.
# Rows past either are written without the template's dropdowns or styles,
# because the template is kept unchanged (see the module docstring).
_FIRST_PLAN_ROW = 5
_TEMPLATE_LAST_PLAN_ROW = 138


# The template's severity scale, exactly as Reference!A19:A23 spells it
# (two spaces after the digit). The Summary counts match on this text.
_SEV_CRITICAL = "4  Critical"
_SEV_MAJOR = "3  Major"
_SEV_MINOR = "2  Minor"
_SEV_BEST_PRACTICE = "1  Best practice"
# Automated impact -> the template's scale. Conservative on purpose: the
# template's levels judge access to core and secondary tasks, which a scan
# cannot know, so no impact maps to "5  Blocker" and a person raises it.
_SEVERITY_FROM_IMPACT = {
    "critical": _SEV_CRITICAL,
    "serious": _SEV_MAJOR,
    "major": _SEV_MAJOR,
    "moderate": _SEV_MINOR,
    "minor": _SEV_MINOR,
    "info": _SEV_MINOR,
}

# The template's starting value for Verified by a11y team (column K).
_NOT_TESTED_BY_TEAM = "Not tested by UM A11y team"

# Criteria the ITS Reference sheet does not list (it predates WCAG 2.2 and
# covers A/AA only), placed in the category of the listed criterion they
# sit beside, so the Category dropdown and the Summary counts still apply.
_EXTRA_CATEGORY = {
    # WCAG 2.2 A/AA additions.
    "2.4.11": "Keyboard and other Input",
    "2.5.7": "Keyboard and other Input",
    "2.5.8": "Keyboard and other Input",
    "3.2.6": "Navigation and Consistency",
    "3.3.7": "Instructions",
    "3.3.8": "Instructions",
    # AAA, for scans checked at Level AAA.
    "1.2.6": "Text Alternatives",
    "1.2.7": "Text Alternatives",
    "1.2.8": "Text Alternatives",
    "1.2.9": "Text Alternatives",
    "1.3.6": "Semantics and Structure",
    "1.4.6": "Contrast",
    "1.4.7": "Media Control",
    "1.4.8": "Adaptability",
    "1.4.9": "Text Alternatives",
    "2.1.3": "Keyboard and other Input",
    "2.2.3": "User Controls",
    "2.2.4": "User Controls",
    "2.2.5": "User Controls",
    "2.2.6": "User Controls",
    "2.3.2": "Media Control",
    "2.3.3": "Media Control",
    "2.4.8": "Navigation and Consistency",
    "2.4.9": "Semantics and Structure",
    "2.4.10": "Semantics and Structure",
    "2.4.12": "Keyboard and other Input",
    "2.4.13": "Keyboard and other Input",
    "2.5.5": "Keyboard and other Input",
    "2.5.6": "Keyboard and other Input",
    "3.1.3": "Language",
    "3.1.4": "Language",
    "3.1.5": "Language",
    "3.1.6": "Language",
    "3.2.5": "User Controls",
    "3.3.5": "Instructions",
    "3.3.6": "Instructions",
    "3.3.9": "Instructions",
}

# The user groups in the template's own words (Reference Impact column B),
# for criteria that sheet does not list.
_ABILITY_GROUPS = {
    "vision": "Blind, Low vision",
    "hearing": "Deaf, Hard of hearing",
    "motor": "Mobility disability",
    "cognition": "Neurodivergent",
}

# Each check by its interface name. Mirrors CHECK_LABEL in
# web/frontend/src/lib/terms.ts so the workbook names a check as the app does.
_CHECK_LABEL = {
    "axe": "Rule check (axe)",
    "alfa": "Rule check (Alfa)",
    "keyboard": "Keyboard check",
    "responsive": "Zoom and layout check",
    "focus": "Focus check",
    "visual": "Motion and reading-order check",
    "semantic": "AI review",
    "image": "Image text check",
}

# Occurrence statuses as the app shows them (STATUS_LABEL in terms.ts).
_STATUS_LABEL = {
    "new": "New",
    "reviewing": "Reviewing",
    "in_progress": "In progress",
    "remediated": "Fixed",
    "accepted_risk": "Accepted risk",
    "false_positive": "Not a problem",
}

# How sure Axcess is, in the words of the app's How sure column, with what
# that asks of the reader.
_LANE_NOTE = {
    "likely_barrier": (
        f"How sure: {REVIEW_LANE_LABELS['likely_barrier']}. A rule check failed. "
        "Check it on the page to confirm."
    ),
    "expert_review": (
        f"How sure: {REVIEW_LANE_LABELS['expert_review']}. A person must decide "
        "whether this is a real problem."
    ),
}

# The template writes each criterion as "<number> <en dash> <name>".
_EN_DASH = "\u2013"
_SC_PREFIX = re.compile(r"^\s*(\d+\.\d+\.\d+)")
_UNDERSTANDING_URL = "https://www.w3.org/WAI/WCAG22/Understanding/{slug}.html"


@dataclass(frozen=True)
class _Reference:
    """What the template's own reference sheets say about each criterion."""

    #: SC number -> the criterion as the template writes it ("1.4.3", an en dash,
    #: "Contrast (Minimum)").
    criterion: dict[str, str]
    #: SC number -> the template's category for it.
    category: dict[str, str]
    #: SC number -> the template's "Most affects people who are/have" text.
    affects: dict[str, str]
    #: SC number -> the template's Impact sentence.
    impact: dict[str, str]


@dataclass(frozen=True)
class _Issue:
    """One Axcess issue and the occurrences it contributes as rows."""

    row: issues.IssueRow
    #: Every occurrence not marked "Not a problem", in page order.
    occurrences: list[IssueLocation]


def _sc_of(text: Any) -> str | None:
    match = _SC_PREFIX.match(str(text or ""))
    return match.group(1) if match else None


def _clean(value: str) -> str:
    """Drop the control characters a worksheet cell cannot hold."""
    return str(ILLEGAL_CHARACTERS_RE.sub("", value))


def _put_text(ws: Worksheet, coordinate: str, value: str | None) -> None:
    """Write ``value`` as plain text, never as a formula.

    openpyxl turns any string that starts with ``=`` into a formula. Scan
    text (a page title, a URL, an element locator) comes from the crawled
    site, so it is forced back to text.
    """
    cell = ws[coordinate]
    cell.value = _clean(value) if value is not None else None
    if cell.data_type == "f":
        cell.data_type = "s"


def _put_link(ws: Worksheet, coordinate: str, text: str, url: str) -> None:
    """Write ``text`` linked to ``url``, in the cell's own template style."""
    _put_text(ws, coordinate, text)
    ws[coordinate].hyperlink = url


def _scan_date(scan: ExportScan) -> datetime | None:
    raw = scan.finished_at or scan.started_at
    if not raw:
        return None
    try:
        parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
    except (ValueError, TypeError):
        return None
    return datetime(parsed.year, parsed.month, parsed.day)


def _human_date(when: datetime | None) -> str:
    return when.strftime("%B %-d, %Y") if when else "an unknown date"


def _plural(count: int, one: str, many: str) -> str:
    return f"{count:,} {one if count == 1 else many}"


def _read_reference(wb: Workbook) -> _Reference:
    """Index the template's Reference and Reference Impact sheets by SC number."""
    criterion: dict[str, str] = {}
    category: dict[str, str] = {}
    ref = wb[_REFERENCE]
    # Row 4 names each category (B4:K4); the criteria sit below it.
    for col in range(2, 12):
        name = ref.cell(row=4, column=col).value
        if not name:
            continue
        for row in range(5, 17):
            text = ref.cell(row=row, column=col).value
            sc = _sc_of(text)
            if sc:
                criterion[sc] = str(text).strip()
                category[sc] = str(name).strip()
    affects: dict[str, str] = {}
    impact: dict[str, str] = {}
    for text, who, why in wb[_IMPACT].iter_rows(min_col=1, max_col=3, values_only=True):
        sc = _sc_of(text)
        if not sc:
            continue
        criterion.setdefault(sc, str(text).strip())
        if who:
            affects[sc] = str(who).strip()
        if why:
            impact[sc] = str(why).strip()
    return _Reference(criterion=criterion, category=category, affects=affects, impact=impact)


def _collect_issues(
    conn: sqlite3.Connection, scan_id: int, page_order: dict[str, int]
) -> list[_Issue]:
    """Every issue that is a barrier, with each of its occurrences.

    The template asks for barriers only, so issues Axcess marks "For
    information" and occurrences someone marked "Not a problem" are dropped.
    """
    collected: list[_Issue] = []
    for row in issues.list_issues(conn, scan_id):
        if row.review_lane == "informational":
            continue
        locations, _ = issue_locations(conn, row, limit=None)
        kept = [loc for loc in locations if loc.status != "false_positive"]
        if not kept:
            continue
        kept.sort(
            key=lambda loc: (
                page_order.get(loc.page_url, len(page_order)),
                loc.finding_id or 0,
                loc.image_url or "",
            )
        )
        collected.append(_Issue(row=row, occurrences=kept))
    return collected


def _scan_pages(conn: sqlite3.Connection, scan_id: int) -> list[tuple[str, str | None]]:
    rows = conn.execute(
        "SELECT url_normalized, title FROM pages WHERE scan_id = ? ORDER BY id",
        (scan_id,),
    ).fetchall()
    return [(str(r[0]), r[1]) for r in rows]


def _test_names(pages: list[tuple[str, str | None]]) -> dict[str, str]:
    """A unique Test name per page URL: its title, told apart by path if needed."""
    taken: set[str] = set()
    names: dict[str, str] = {}
    for url, title in pages:
        path = urlsplit(url).path or "/"
        base = " ".join(str(title or "").split()) or path
        name = base
        if name in taken:
            name = f"{base} ({path})"
        suffix = 2
        while name in taken:
            name = f"{base} ({path}) {suffix}"
            suffix += 1
        taken.add(name)
        names[url] = name
    return names


def _criterion_text(row: issues.IssueRow, ref: _Reference) -> str:
    if not row.wcag_sc:
        return "No WCAG criterion (best practice)"
    if row.wcag_sc in ref.criterion:
        return ref.criterion[row.wcag_sc]
    if row.wcag_name:
        return f"{row.wcag_sc} {_EN_DASH} {row.wcag_name}"
    return row.wcag_sc


def _info_link(row: issues.IssueRow, criterion_names: dict[str, str]) -> tuple[str, str] | None:
    """The template's Info column: W3C's page on the criterion, else the rule's page.

    W3C names each Understanding page after the criterion's official name, so
    the link is built only from the coverage matrix's official names.
    """
    if row.wcag_sc and row.wcag_sc in criterion_names:
        slug = re.sub(r"[^a-z0-9]+", "-", criterion_names[row.wcag_sc].lower()).strip("-")
        return f"{row.wcag_sc} Info", _UNDERSTANDING_URL.format(slug=slug)
    if row.help_url:
        return "Rule info", row.help_url
    return None


def _severity(row: issues.IssueRow) -> str:
    if row.conformance == "BP" or not row.wcag_sc:
        return _SEV_BEST_PRACTICE
    return _SEVERITY_FROM_IMPACT.get((row.impact or "").lower(), _SEV_MINOR)


def _notes(issue: _Issue, number: int, loc: IssueLocation, scan_date: str) -> str:
    """One occurrence's Issue notes: what is wrong, where, how sure, how to fix it."""
    row = issue.row
    total = len(issue.occurrences)
    lines = [
        f"{scan_date}, Axcess scan: {plain_text(row.title)} (occurrence {number:,} of {total:,})"
    ]
    if row.description:
        lines.append(plain_text(row.description))

    page = loc.page_url or "No page recorded"
    if loc.page_title:
        page = f"{' '.join(loc.page_title.split())}: {page}"
    lines.append(f"Page: {page}")
    if loc.selector:
        lines.append(f"Element locator (CSS selector): {loc.selector}")
    elif loc.image_url:
        lines.append(f"Image: {loc.image_url}")
    if loc.detail:
        lines.append(f"What the check found: {plain_text(loc.detail)}")
    if loc.revealed_by:
        lines.append(f'To see it, first open "{loc.revealed_by}".')

    found_by = _CHECK_LABEL.get(row.pipeline, row.pipeline)
    lane = _LANE_NOTE.get(row.review_lane, "")
    lines.append(f"Found by: {found_by}. {lane}".strip())
    if row.repeat_page_count:
        lines.append(
            "Some elements with this issue also appear, unchanged, on "
            f"{_plural(row.repeat_page_count, 'more page', 'more pages')}, so they are "
            "likely part of a shared header, menu or footer. Axcess lists each one once, "
            "on the first page it was found."
        )

    steps = [plain_text(step) for step in row.fix_steps if step and step.strip()]
    if steps:
        lines.append("How to fix:")
        lines.extend(f"• {step}" for step in steps)
    if loc.status and loc.status != "new":
        lines.append(f"Status in Axcess: {_STATUS_LABEL.get(loc.status, loc.status)}.")
    return "\n".join(lines)


def _fill_coversheet(
    ws: Worksheet,
    scan: ExportScan,
    scan_day: datetime | None,
    evaluation_record: dict[str, Any],
    issues_found: list[_Issue],
) -> None:
    if scan_day is not None:
        ws["B3"].value = scan_day
    _put_text(ws, "B8", urlsplit(scan.seed_url).hostname or scan.seed_url)
    purpose = str(evaluation_record.get("purpose") or "").strip()
    if evaluation_record.get("exists") and purpose:
        _put_text(ws, "B13", purpose)
    _put_link(ws, "B17", scan.seed_url, scan.seed_url)
    # Login info (B18) is left for the team. A sign-in scan is an ordinary
    # report with no record that it signed in, so the scan cannot say.

    checks: list[str] = []
    if scan.axe_pages_scanned:
        checks.append(_CHECK_LABEL["axe"])
    if scan.alfa_pages_scanned:
        checks.append(_CHECK_LABEL["alfa"])
    for issue in issues_found:
        label = _CHECK_LABEL.get(issue.row.pipeline)
        if label and label not in checks:
            checks.append(label)
    tools = (
        f"Axcess automated scan, WCAG {scan.wcag_version} Level {scan.axe_level}, "
        "in a Chromium browser."
    )
    if checks:
        tools += " Checks with results: " + ", ".join(checks) + "."
    tools += "\nAdd the screen readers, browsers and tools the team used for manual testing."
    _put_text(ws, "B19", tools)


def _fill_test_plan(
    ws: Worksheet,
    pages: list[tuple[str, str | None]],
    names: dict[str, str],
) -> None:
    """One row per scanned page: the Test name and its URL.

    The Description column is left empty for the team to write.
    """
    for row in range(_FIRST_PLAN_ROW, _TEMPLATE_LAST_PLAN_ROW + 1):
        for col in "ABCD":
            ws[f"{col}{row}"].value = None
    for offset, (url, _title) in enumerate(pages):
        row = _FIRST_PLAN_ROW + offset
        _put_text(ws, f"A{row}", names[url])
        _put_link(ws, f"B{row}", url, url)


def _fill_findings(
    ws: Worksheet,
    issues_found: list[_Issue],
    ref: _Reference,
    names: dict[str, str],
    scan_date: str,
) -> None:
    """Write one row per occurrence, starting at row 2."""
    total = sum(len(issue.occurrences) for issue in issues_found)
    last = max(_TEMPLATE_LAST_ISSUE_ROW, _FIRST_ISSUE_ROW + total - 1)
    criterion_names = {c.sc: c.name for c in coverage_matrix.load_matrix()}
    for row_number in range(_FIRST_ISSUE_ROW, last + 1):
        for col in "BCDEFGHI":
            ws[f"{col}{row_number}"].value = None
            ws[f"{col}{row_number}"].hyperlink = None
        ws[f"A{row_number}"].value = row_number - _FIRST_ISSUE_ROW + 1
        ws[f"J{row_number}"].value = False
        ws[f"K{row_number}"].value = _NOT_TESTED_BY_TEAM

    r = _FIRST_ISSUE_ROW
    for issue in issues_found:
        row = issue.row
        sc = row.wcag_sc or ""
        # The same for every occurrence of the issue.
        category = ref.category.get(sc) or _EXTRA_CATEGORY.get(sc)
        criterion = _criterion_text(row, ref)
        info = _info_link(row, criterion_names)
        severity = _severity(row)
        why = ref.impact.get(sc) or (plain_text(row.why_matters) if row.why_matters else None)
        who = ref.affects.get(sc) or ", ".join(
            _ABILITY_GROUPS[a] for a in row.abilities_affected if a in _ABILITY_GROUPS
        )
        for number, loc in enumerate(issue.occurrences, start=1):
            _put_text(ws, f"B{r}", names.get(loc.page_url))
            _put_text(ws, f"C{r}", _notes(issue, number, loc, scan_date))
            _put_text(ws, f"D{r}", category)
            _put_text(ws, f"E{r}", criterion)
            if info:
                _put_link(ws, f"F{r}", *info)
            ws[f"G{r}"].value = severity
            _put_text(ws, f"H{r}", why)
            _put_text(ws, f"I{r}", who or None)
            r += 1


def _fill_summary(
    ws: Worksheet,
    scan: ExportScan,
    issues_found: list[_Issue],
    scan_date: str,
) -> None:
    """Write the Narrative Summary cell. The template's formulas are left as they are."""
    sure = sum(1 for issue in issues_found if issue.row.review_lane == "likely_barrier")
    unsure = len(issues_found) - sure
    rows = sum(len(issue.occurrences) for issue in issues_found)
    template_rows = _TEMPLATE_LAST_ISSUE_ROW - _FIRST_ISSUE_ROW + 1
    beyond = rows - template_rows
    # The template's counts stop at its last row; say so rather than let a
    # reader trust a total that leaves rows out.
    counted = (
        f"\nThe counts below include the first {template_rows} rows only. "
        f"{_plural(beyond, 'more row is', 'more rows are')} listed on Findings Issues "
        f"after row {_TEMPLATE_LAST_ISSUE_ROW}."
        if beyond > 0
        else ""
    )
    _put_text(
        ws,
        "A4",
        f"Axcess scanned {_plural(scan.page_count, 'page', 'pages')} of {scan.seed_url} "
        f"on {scan_date}. It found {_plural(len(issues_found), 'issue', 'issues')} to review: "
        f"{sure:,} where a rule check failed, and {unsure:,} that a person must confirm. "
        f"Findings Issues has one row for each place an issue appears (an occurrence): "
        f"{_plural(rows, 'row', 'rows')} in all. "
        "Automated checks find only some accessibility problems, so test by hand as well.\n"
        "Severity is a first estimate from the automated impact rating. No issue is marked "
        "Blocker, because that depends on which tasks are core to the site, which a person "
        "must judge. Confirm each issue and its severity before you share this report." + counted,
    )


def render_its_review(scan: ExportScan, *, conn: sqlite3.Connection) -> bytes:
    """Fill the U-M ITS Technical Review template from ``scan``; return .xlsx bytes.

    ``conn`` is required: the issue rows come from the same
    ``issues.list_issues()`` projection the app's Issues table shows.
    """
    source = resources.files(_TEMPLATE_PACKAGE).joinpath(*_TEMPLATE_PATH)
    wb = load_workbook(io.BytesIO(source.read_bytes()))
    ref = _read_reference(wb)

    scan_day = _scan_date(scan)
    scan_date = _human_date(scan_day)
    pages = _scan_pages(conn, scan.id)
    names = _test_names(pages)
    issues_found = _collect_issues(
        conn, scan.id, {url: index for index, (url, _) in enumerate(pages)}
    )

    _fill_coversheet(
        wb[_COVER], scan, scan_day, evaluation.get_evaluation(conn, scan.id), issues_found
    )
    _fill_test_plan(wb[_PLAN], pages, names)
    _fill_findings(wb[_FINDINGS], issues_found, ref, names, scan_date)
    _fill_summary(wb[_SUMMARY], scan, issues_found, scan_date)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
