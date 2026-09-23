"""Semantic fingerprint of an .xlsx workbook, for export-equivalence checks.

An .xlsx file is a zip whose bytes change on every save (member timestamps,
``docProps/core.xml`` created/modified), so two renders of the same workbook
never compare equal byte for byte. This module parses the workbook back with
openpyxl and reduces it to a JSON-serializable dict that holds everything a
reader of the workbook can observe and nothing that varies between saves:

* per sheet, in workbook order: title, sheet state, whether it is the active
  sheet, tab color and gridline visibility;
* every non-empty cell (a value, or a non-default style) with its value,
  data type, number format and a style reference;
* hyperlinks (target, location, tooltip, display), merged ranges, freeze
  panes, the auto-filter range, data validations, embedded images, column
  widths and row heights;
* workbook-level defined names and the named-style list.

Cell styles are interned: each distinct font/fill/border/alignment/protection
and named-style combination is stored once under a short content hash, and
cells refer to that hash. An unrelated style change therefore does not
renumber every other cell, and a diff points at the cells that changed.

The same fingerprint backs the unit-test goldens (``tests/unit/golden/*.xlsx.json``)
and ``scripts/export_diff.py``. It lives under ``tests/support`` so both can
import it without making ``tests`` a package: pytest already puts ``tests/``
on ``sys.path`` through the root ``conftest.py``, and the script adds it
explicitly.
"""

from __future__ import annotations

import difflib
import hashlib
import io
import json
import math
from datetime import date, datetime, time, timedelta
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

#: Bumped whenever the fingerprint shape changes, so a golden recorded by an
#: older helper fails loudly instead of comparing unlike structures.
FINGERPRINT_VERSION = 1

# One pixel is 9525 EMU in DrawingML; openpyxl anchors images in EMU.
_EMU_PER_PIXEL = 9525


def fingerprint_xlsx(data: bytes) -> dict[str, Any]:
    """Return the semantic fingerprint of the workbook in ``data``."""
    workbook = load_workbook(io.BytesIO(data))
    styles: dict[str, Any] = {}
    active = workbook.active
    sheets = [_sheet(ws, is_active=ws is active, styles=styles) for ws in workbook.worksheets]
    return {
        "fingerprint_version": FINGERPRINT_VERSION,
        "sheet_names": [ws.title for ws in workbook.worksheets],
        "defined_names": _defined_names(workbook.defined_names),
        "named_styles": sorted(_named_style_names(workbook)),
        "sheets": sheets,
        # Keys are written sorted; this one sorts ahead of "sheets", so a
        # changed style definition leads a (possibly truncated) diff instead
        # of trailing hundreds of changed cell references.
        "cell_styles": dict(sorted(styles.items())),
    }


def dumps_fingerprint(fingerprint: dict[str, Any]) -> str:
    """Serialize a fingerprint as stable, line-oriented JSON.

    Containers that hold only scalars (a cell, a style part, a merged-range
    list) are written on one line; everything else is indented. The result
    is valid JSON that parses back to ``fingerprint``, and a changed cell
    shows up as one changed line in a diff.
    """
    return _emit(fingerprint, 0) + "\n"


def fingerprint_diff(
    expected: dict[str, Any],
    actual: dict[str, Any],
    *,
    labels: tuple[str, str] = ("expected", "actual"),
    limit: int = 80,
) -> list[str]:
    """Unified-diff lines between two fingerprints, capped at ``limit`` lines.

    The first line names what changed (top-level parts and sheet titles), so
    a truncated diff still says where to look.
    """
    if expected == actual:
        return []
    changed = [key for key in sorted(set(expected) | set(actual)) if key != "sheets"]
    changed = [key for key in changed if expected.get(key) != actual.get(key)]
    expected_sheets = {sheet["title"]: sheet for sheet in expected.get("sheets", [])}
    actual_sheets = {sheet["title"]: sheet for sheet in actual.get("sheets", [])}
    changed += [
        f"sheet {title!r}"
        for title in [*expected_sheets, *(t for t in actual_sheets if t not in expected_sheets)]
        if expected_sheets.get(title) != actual_sheets.get(title)
    ]
    summary = "changed: " + ", ".join(changed)
    lines = list(
        difflib.unified_diff(
            dumps_fingerprint(expected).splitlines(),
            dumps_fingerprint(actual).splitlines(),
            fromfile=labels[0],
            tofile=labels[1],
            lineterm="",
            n=2,
        )
    )
    if len(lines) > limit:
        hidden = len(lines) - limit
        lines = [*lines[:limit], f"... {hidden} more diff line(s) not shown"]
    return [summary, *lines]


# --------------------------------------------------------------------------
# Sheets.
# --------------------------------------------------------------------------


def _sheet(ws: Any, *, is_active: bool, styles: dict[str, Any]) -> dict[str, Any]:
    cells: list[list[Any]] = []
    hyperlinks: list[list[Any]] = []
    for _position, cell in sorted(ws._cells.items()):
        link = cell.hyperlink
        if link is not None:
            hyperlinks.append(
                [
                    cell.coordinate,
                    link.target,
                    link.location,
                    link.tooltip,
                    link.display,
                ]
            )
        if cell.value is None and not cell.has_style:
            continue
        cells.append(
            [
                cell.coordinate,
                _value(cell.value),
                cell.data_type,
                cell.number_format,
                _style_key(cell, styles),
            ]
        )
    tab_color = ws.sheet_properties.tabColor
    return {
        "title": ws.title,
        "state": ws.sheet_state,
        "active": is_active,
        "tab_color": _color(tab_color),
        "show_grid_lines": ws.sheet_view.showGridLines,
        "freeze_panes": ws.freeze_panes,
        "auto_filter": ws.auto_filter.ref,
        "merged": [
            str(merged)
            for merged in sorted(
                ws.merged_cells.ranges,
                key=lambda r: (r.min_row, r.min_col, r.max_row, r.max_col),
            )
        ],
        "columns": _columns(ws),
        "rows": _rows(ws),
        "data_validations": _data_validations(ws),
        "images": _images(ws),
        "hyperlinks": hyperlinks,
        "cells": cells,
    }


def _columns(ws: Any) -> dict[str, list[Any]]:
    """Width and visibility per column, with openpyxl's range grouping undone.

    openpyxl writes adjacent columns with identical settings as one ``<col>``
    range and reads them back under the first letter only. Expanding the
    range keeps "C..F are 14 wide" identical however the writer grouped it.
    """
    widths: dict[int, list[Any]] = {}
    for key, dim in ws.column_dimensions.items():
        first = dim.min or column_index_from_string(key)
        last = dim.max or first
        for index in range(first, last + 1):
            widths[index] = [dim.width, bool(dim.hidden)]
    return {get_column_letter(index): value for index, value in sorted(widths.items())}


def _rows(ws: Any) -> dict[str, list[Any]]:
    out: dict[str, list[Any]] = {}
    for index, dim in sorted(ws.row_dimensions.items()):
        if dim.height is None and not dim.hidden:
            continue
        out[str(index)] = [dim.height, bool(dim.hidden)]
    return out


def _data_validations(ws: Any) -> list[dict[str, Any]]:
    rules = []
    for dv in ws.data_validations.dataValidation:
        rules.append(
            {
                "type": dv.type,
                "operator": dv.operator,
                "formula1": dv.formula1,
                "formula2": dv.formula2,
                "sqref": " ".join(sorted(str(ref) for ref in dv.sqref.ranges)),
                "allow_blank": bool(dv.allow_blank),
                "show_dropdown": bool(dv.showDropDown),
                "show_input_message": bool(dv.showInputMessage),
                "show_error_message": bool(dv.showErrorMessage),
                "prompt_title": dv.promptTitle,
                "prompt": dv.prompt,
                "error_title": dv.errorTitle,
                "error": dv.error,
                "error_style": dv.errorStyle,
            }
        )
    # Validation order in the XML is not observable in Excel.
    return sorted(rules, key=lambda rule: json.dumps(rule, sort_keys=True))


def _images(ws: Any) -> list[dict[str, Any]]:
    images = []
    for image in ws._images:
        anchor = image.anchor
        entry: dict[str, Any] = {"anchor_type": type(anchor).__name__}
        if isinstance(anchor, str):
            entry["anchor"] = anchor
            entry["offset"] = [0, 0]
            width, height = image.width, image.height
        else:
            marker = getattr(anchor, "_from", None)
            if marker is not None:
                entry["anchor"] = f"{get_column_letter(marker.col + 1)}{marker.row + 1}"
                entry["offset"] = [marker.colOff, marker.rowOff]
            to = getattr(anchor, "to", None)
            if to is not None:
                entry["to"] = f"{get_column_letter(to.col + 1)}{to.row + 1}"
                entry["to_offset"] = [to.colOff, to.rowOff]
            ext = getattr(anchor, "ext", None)
            if ext is not None and getattr(ext, "width", None) is not None:
                # The displayed size lives in the anchor; ``image.width`` on a
                # loaded workbook is the source bitmap's native size.
                width = round(ext.width / _EMU_PER_PIXEL, 2)
                height = round(ext.height / _EMU_PER_PIXEL, 2)
            else:
                width, height = image.width, image.height
        entry["width"] = width
        entry["height"] = height
        entry["native_size"] = [image.width, image.height]
        entry["format"] = getattr(image, "format", None)
        entry["sha256"] = _image_digest(image)
        images.append(entry)
    return sorted(images, key=lambda item: json.dumps(item, sort_keys=True))


def _image_digest(image: Any) -> str | None:
    try:
        return hashlib.sha256(image._data()).hexdigest()[:16]
    except Exception:  # an unreadable image is still reported by its anchor
        return None


# --------------------------------------------------------------------------
# Workbook-level parts.
# --------------------------------------------------------------------------


def _defined_names(names: Any) -> dict[str, dict[str, Any]]:
    out: dict[str, dict[str, Any]] = {}
    for name, defined in sorted(names.items()):
        out[name] = {
            "value": defined.attr_text,
            "local_sheet_id": defined.localSheetId,
            "hidden": bool(defined.hidden),
        }
    return out


def _named_style_names(workbook: Any) -> list[str]:
    return [str(getattr(style, "name", style)) for style in workbook.named_styles]


# --------------------------------------------------------------------------
# Values and styles.
# --------------------------------------------------------------------------


def _value(value: Any) -> Any:
    """A JSON-safe value that keeps int / float / str / bool distinct."""
    if value is None or isinstance(value, bool | int | str):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else {"float": repr(value)}
    if isinstance(value, datetime | date | time):
        return {type(value).__name__: value.isoformat()}
    if isinstance(value, timedelta):
        return {"timedelta": value.total_seconds()}
    return {type(value).__name__: str(value)}


def _style_key(cell: Any, styles: dict[str, Any]) -> str:
    style = {
        "named_style": cell.style,
        "font": _font(cell.font),
        "fill": _fill(cell.fill),
        "border": _border(cell.border),
        "alignment": _alignment(cell.alignment),
        "protection": {"locked": cell.protection.locked, "hidden": cell.protection.hidden},
    }
    encoded = json.dumps(style, sort_keys=True, separators=(",", ":"))
    key = hashlib.sha256(encoded.encode("utf-8")).hexdigest()[:12]
    styles.setdefault(key, style)
    return key


def _color(color: Any) -> dict[str, Any] | None:
    if color is None:
        return None
    return {
        "type": color.type,
        "value": _value(color.value),
        "tint": color.tint,
    }


def _font(font: Any) -> dict[str, Any]:
    return {
        "name": font.name,
        "size": font.size,
        "bold": font.bold,
        "italic": font.italic,
        "underline": font.underline,
        "strike": font.strike,
        "vert_align": font.vertAlign,
        "color": _color(font.color),
        "scheme": font.scheme,
        "family": font.family,
        "charset": font.charset,
        "outline": font.outline,
        "shadow": font.shadow,
        "condense": font.condense,
        "extend": font.extend,
    }


def _fill(fill: Any) -> dict[str, Any]:
    if getattr(fill, "tagname", "") == "gradientFill":
        return {
            "kind": "gradient",
            "type": fill.type,
            "degree": fill.degree,
            "left": fill.left,
            "right": fill.right,
            "top": fill.top,
            "bottom": fill.bottom,
            "stops": [[stop.position, _color(stop.color)] for stop in fill.stop],
        }
    return {
        "kind": "pattern",
        "pattern": fill.fill_type,
        "fg": _color(fill.fgColor),
        "bg": _color(fill.bgColor),
    }


def _side(side: Any) -> dict[str, Any] | None:
    if side is None:
        return None
    return {"style": side.style, "color": _color(side.color)}


def _border(border: Any) -> dict[str, Any]:
    return {
        "left": _side(border.left),
        "right": _side(border.right),
        "top": _side(border.top),
        "bottom": _side(border.bottom),
        "diagonal": _side(border.diagonal),
        "vertical": _side(border.vertical),
        "horizontal": _side(border.horizontal),
        "diagonal_up": border.diagonalUp,
        "diagonal_down": border.diagonalDown,
        "outline": border.outline,
    }


def _alignment(alignment: Any) -> dict[str, Any]:
    return {
        "horizontal": alignment.horizontal,
        "vertical": alignment.vertical,
        "wrap_text": alignment.wrap_text,
        "shrink_to_fit": alignment.shrink_to_fit,
        "indent": alignment.indent,
        "text_rotation": alignment.text_rotation,
        "relative_indent": alignment.relativeIndent,
        "justify_last_line": alignment.justifyLastLine,
        "reading_order": alignment.readingOrder,
    }


# --------------------------------------------------------------------------
# Serialization.
# --------------------------------------------------------------------------


def _is_flat(value: dict[str, Any] | list[Any]) -> bool:
    items = value.values() if isinstance(value, dict) else value
    return all(not isinstance(item, dict | list) for item in items)


def _emit(value: Any, level: int) -> str:
    if not isinstance(value, dict | list) or not value or _is_flat(value):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    pad = "  " * (level + 1)
    close = "  " * level
    if isinstance(value, dict):
        body = ",\n".join(
            f"{pad}{json.dumps(key, ensure_ascii=False)}: {_emit(item, level + 1)}"
            for key, item in sorted(value.items())
        )
        return "{\n" + body + "\n" + close + "}"
    body = ",\n".join(f"{pad}{_emit(item, level + 1)}" for item in value)
    return "[\n" + body + "\n" + close + "]"
