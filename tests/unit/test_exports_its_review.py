"""The U-M ITS technical review: the team's own template, filled in from one scan."""

from __future__ import annotations

import io
import sqlite3
from datetime import datetime
from importlib import resources
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from support.rich_scan import dom_finding, seed_rich_reviews, seed_rich_scan
from test_audit_report import _scan_with_real_findings  # tests/unit on sys.path
from typer.testing import CliRunner

from audit import cli
from audit.config import Settings
from audit.db import repo
from audit.exports.collector import collect_scan
from audit.exports.its_review import render_its_review
from audit.occurrence_reviews import Decision, save_review
from audit.web import issues

_HEADERS = {
    "A": "number",
    "B": "test",
    "C": "notes",
    "D": "category",
    "E": "criterion",
    "F": "info",
    "G": "severity",
    "H": "impact",
    "I": "affects",
    "J": "vendor_ready",
    "K": "verified",
}
_SEVERITIES = {"4  Critical", "3  Major", "2  Minor", "1  Best practice"}


def _template() -> Workbook:
    source = resources.files("audit.exports").joinpath("templates", "its_technical_review.xlsx")
    return load_workbook(io.BytesIO(source.read_bytes()))


def _render(conn: sqlite3.Connection, scan_id: int) -> Workbook:
    out = render_its_review(collect_scan(conn, scan_id), conn=conn)
    return load_workbook(io.BytesIO(out))


def _issue_rows(wb: Workbook) -> list[dict[str, Any]]:
    ws = wb["Findings Issues"]
    rows = []
    for r in range(2, ws.max_row + 1):
        if ws[f"G{r}"].value is None:
            continue
        row = {key: ws[f"{col}{r}"].value for col, key in _HEADERS.items()}
        row["info_link"] = ws[f"F{r}"].hyperlink.target if ws[f"F{r}"].hyperlink else None
        rows.append(row)
    return rows


def _all_text(wb: Workbook) -> str:
    return "\n".join(
        str(cell.value)
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None
    )


def _rich(conn: sqlite3.Connection) -> int:
    """The rich scan with a person's reviews (see ``seed_rich_reviews``)."""
    scan_id = seed_rich_scan(conn)
    seed_rich_reviews(conn, scan_id)
    return scan_id


def _review(
    conn: sqlite3.Connection,
    scan_id: int,
    finding_ids: list[int],
    *,
    decision: Decision | None = "report",
    notes: str = "",
) -> None:
    page_id = conn.execute(
        "SELECT page_id FROM page_a11y_findings WHERE id = ?", (finding_ids[0],)
    ).fetchone()[0]
    save_review(
        conn,
        scan_id=scan_id,
        page_id=int(page_id),
        a11y_finding_ids=finding_ids,
        page_image_ids=[],
        decision=decision,
        notes=notes,
    )


def _finding_id(conn: sqlite3.Connection, scan_id: int, target: str) -> int:
    return int(
        conn.execute(
            "SELECT id FROM page_a11y_findings WHERE scan_id = ? AND target_selector = ? "
            "ORDER BY id DESC LIMIT 1",
            (scan_id, target),
        ).fetchone()[0]
    )


def _title(conn: sqlite3.Connection, scan_id: int, issue_key: str) -> str:
    (row,) = [r for r in issues.list_issues(conn, scan_id) if r.issue_key == issue_key]
    return row.title


def _tests(wb: Workbook) -> list[tuple[str, str, str | None]]:
    """Test Plan rows: (Test, URL or location, its link)."""
    plan = wb["Test Plan"]
    return [
        (
            str(plan[f"A{r}"].value),
            str(plan[f"B{r}"].value),
            plan[f"B{r}"].hyperlink.target if plan[f"B{r}"].hyperlink else None,
        )
        for r in range(5, plan.max_row + 1)
        if plan[f"A{r}"].value
    ]


def test_keeps_the_templates_sheets_and_reference_tables(tmp_db: sqlite3.Connection) -> None:
    scan_id = seed_rich_scan(tmp_db)
    wb = _render(tmp_db, scan_id)
    template = _template()

    assert wb.sheetnames == template.sheetnames
    for name in ("Reference", "Reference Impact"):
        got = [[c.value for c in row] for row in wb[name].iter_rows()]
        want = [[c.value for c in row] for row in template[name].iter_rows()]
        assert got == want, f"{name} must stay exactly as the ITS team wrote it"
    # The dropdowns still read the template's named lists.
    for name in ("Severity", "GenericTestList", "TextAlternatives"):
        assert name in wb.defined_names


def test_lists_only_the_elements_a_person_reported(tmp_db: sqlite3.Connection) -> None:
    """Reported elements only: not checked yet, and repeats, are left out."""
    scan_id = _rich(tmp_db)
    wb = _render(tmp_db, scan_id)
    rows = _issue_rows(wb)
    contrast_title = _title(tmp_db, scan_id, "axe:color-contrast")

    # Four reported elements: two contrast, one link, one image.
    assert sorted(r["notes"] for r in rows) == sorted(
        [
            "Muted grey text on white fails contrast.",
            contrast_title,  # reported without notes: what the problem is
            "The first menu link has no name.",
            "The banner's words are not in its alt text.",
        ]
    )
    assert [r["number"] for r in rows] == [1, 2, 3, 4]
    text = _all_text(wb)
    assert "Same muted text as the home page." not in text  # a repeat
    assert "nav > a:nth-child(2)" not in text  # not checked yet


def test_an_issue_row_uses_the_templates_own_words(tmp_db: sqlite3.Connection) -> None:
    scan_id = _rich(tmp_db)
    wb = _render(tmp_db, scan_id)
    impact = {
        str(a).split(" ")[0]: (b, c)
        for a, b, c in _template()["Reference Impact"].iter_rows(max_col=3, values_only=True)
    }

    # The template's own spelling, with an en dash.
    criterion = "1.4.3 \u2013 Contrast (Minimum)"
    contrast = [r for r in _issue_rows(wb) if r["criterion"] == criterion]
    assert len(contrast) == 2
    for row in contrast:
        assert row["test"] == "Home"
        assert row["category"] == "Contrast"
        assert row["severity"] == "3  Major"  # axe "serious"
        assert (row["affects"], row["impact"]) == impact["1.4.3"]
        assert row["info"] == "1.4.3 Info"
        assert row["info_link"] == (
            "https://www.w3.org/WAI/WCAG22/Understanding/contrast-minimum.html"
        )
        assert row["vendor_ready"] is False
        assert row["verified"] == "Not tested by UM A11y team"


def test_one_row_per_element(tmp_db: sqlite3.Connection) -> None:
    """Several occurrences on one element, reviewed together, are one row."""
    scan_id = seed_rich_scan(tmp_db)
    first = _finding_id(tmp_db, scan_id, "p > span.muted")
    # The same element recorded twice (same locator and markup, its own
    # element hash), which the inspector outlines once and reviews together.
    tmp_db.execute(
        "INSERT INTO page_a11y_findings (page_id, scan_id, pipeline, rule_id, wcag_sc, "
        "wcag_level, impact, help, target_selector, failure_summary, html_snippet, "
        "target_hash) SELECT page_id, scan_id, pipeline, rule_id, wcag_sc, wcag_level, "
        "impact, help, target_selector, failure_summary, html_snippet, 'second-record' "
        "FROM page_a11y_findings WHERE id = ?",
        (first,),
    )
    second = _finding_id(tmp_db, scan_id, "p > span.muted")
    assert second != first
    _review(tmp_db, scan_id, [first, second], notes="Low contrast.")

    rows = _issue_rows(_render(tmp_db, scan_id))
    assert [r["notes"] for r in rows] == ["Low contrast."]


def test_test_plan_lists_where_reported_elements_are(tmp_db: sqlite3.Connection) -> None:
    """Pages and page states with reported elements: the clicks first, the URL under."""
    scan_id = seed_rich_scan(tmp_db)
    finding = _finding_id(tmp_db, scan_id, ".details p")
    page_id = tmp_db.execute(
        "SELECT page_id FROM page_a11y_findings WHERE id = ?", (finding,)
    ).fetchone()[0]
    tmp_db.execute(
        "INSERT INTO page_dom_states (page_id, scan_id, state_key, revealed_by, path_labels, "
        "dom) VALUES (?, ?, 'click:more', 'Show more', '[\"Account\", \"Show more\"]', ?)",
        (page_id, scan_id, b"x"),
    )
    tmp_db.execute(
        "UPDATE page_a11y_findings SET revealed_state_key = 'click:more' WHERE id = ?", (finding,)
    )
    _review(tmp_db, scan_id, [finding], notes="Grey text in the opened panel.")
    _review(tmp_db, scan_id, [_finding_id(tmp_db, scan_id, "footer small")], notes="Footer.")

    wb = _render(tmp_db, scan_id)
    assert _tests(wb) == [
        ("Home", "https://example.org/", "https://example.org/"),
        (
            "About us, after clicking “Account” → “Show more”",
            "After clicking “Account” → “Show more”\nhttps://example.org/about",
            "https://example.org/about",
        ),
    ]
    plan = wb["Test Plan"]
    # Only reported places, and the Description column is the team's to write.
    assert all(plan[f"C{r}"].value is None for r in range(5, plan.max_row + 1))
    rows = {r["notes"]: r["test"] for r in _issue_rows(wb)}
    assert rows == {
        "Grey text in the opened panel.": "About us, after clicking “Account” → “Show more”",
        "Footer.": "Home",
    }
    assert wb.defined_names["GenericTestList"].attr_text == "'Test Plan'!$A$5:$A$138"


def test_severity_is_conservative(tmp_db: sqlite3.Connection) -> None:
    scan_id = _rich(tmp_db)
    rows = _issue_rows(_render(tmp_db, scan_id))

    # Blocker needs a person's judgement of which tasks are core.
    assert rows and {r["severity"] for r in rows} <= _SEVERITIES


def test_scan_text_never_becomes_a_formula(tmp_db: sqlite3.Connection) -> None:
    scan_id = seed_rich_scan(tmp_db)
    hostile = '=HYPERLINK("https://attacker.example","Open")'
    page = repo.upsert_page(
        tmp_db,
        scan_id=scan_id,
        url_normalized="https://example.org/hostile",
        status_code=200,
        title=hostile,
        render_mode="js",
        html_hash="8" * 64,
    )
    dom_finding(
        tmp_db,
        scan_id=scan_id,
        page_id=page,
        rule_id="hostile-rule",
        help_text="Hostile page",
        target="=1+1",
        impact="minor",
        wcag_sc="1.3.1",
        wcag_level="A",
    )
    _review(tmp_db, scan_id, [_finding_id(tmp_db, scan_id, "=1+1")], notes="=2+2")
    wb = _render(tmp_db, scan_id)

    cells = [
        cell
        for ws in (wb["Test Plan"], wb["Findings Issues"])
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    # The page title, as the Test and in the Test Plan, and the notes: text.
    assert {c.value for c in cells} == {hostile, "=2+2"}
    assert all(cell.data_type == "s" for cell in cells)


def _style(cell: Any) -> tuple[Any, ...]:
    return (
        repr(cell.font),
        repr(cell.fill),
        repr(cell.border),
        repr(cell.alignment),
        repr(cell.protection),
        cell.number_format,
    )


def test_template_structure_is_unchanged(tmp_db: sqlite3.Connection) -> None:
    """Only data cells change: formulas, lists, names and styles stay as supplied.

    The baseline is the template after one openpyxl load and save, so the
    comparison sees what this exporter changed, not openpyxl's own rewrite.
    That rewrite only normalizes defaults (``shrinkToFit=0`` is dropped) and
    leaves out parts that are empty in the template: six blank drawings and
    an empty list of people.
    """
    scan_id = _rich(tmp_db)
    wb = _render(tmp_db, scan_id)
    buf = io.BytesIO()
    _template().save(buf)
    template = load_workbook(io.BytesIO(buf.getvalue()))

    assert {n: d.attr_text for n, d in wb.defined_names.items()} == {
        n: d.attr_text for n, d in template.defined_names.items()
    }
    for name in template.sheetnames:
        got, want = wb[name], template[name]
        assert [
            (str(v.sqref), v.type, v.formula1) for v in got.data_validations.dataValidation
        ] == [(str(v.sqref), v.type, v.formula1) for v in want.data_validations.dataValidation], (
            name
        )
        assert {str(r) for r in got.merged_cells.ranges} == {
            str(r) for r in want.merged_cells.ranges
        }, name
        assert {k: v.width for k, v in got.column_dimensions.items()} == {
            k: v.width for k, v in want.column_dimensions.items()
        }, name
        assert got.freeze_panes == want.freeze_panes, name
        for row in want.iter_rows():
            for cell in row:
                mine = got[cell.coordinate]
                assert _style(mine) == _style(cell), f"{name}!{cell.coordinate} style"
                if cell.data_type == "f":
                    assert mine.value == cell.value, f"{name}!{cell.coordinate} formula"
    # The template's example value is kept, too.
    assert wb["Coversheet"]["B5"].value == template["Coversheet"]["B5"].value


def test_summary_when_nothing_is_reported_yet(tmp_db: sqlite3.Connection) -> None:
    scan_id = seed_rich_scan(tmp_db)
    wb = _render(tmp_db, scan_id)

    assert _issue_rows(wb) == []
    assert _tests(wb) == []
    narrative = str(wb["Summary"]["A4"].value)
    assert "No element has been checked and reported yet" in narrative
    assert "Checked: a problem to report" in narrative


def test_summary_says_when_its_counts_miss_rows(tmp_db: sqlite3.Connection) -> None:
    scan_id = seed_rich_scan(tmp_db)
    page_id = tmp_db.execute(
        "SELECT id FROM pages WHERE scan_id = ? ORDER BY id LIMIT 1", (scan_id,)
    ).fetchone()[0]
    for n in range(50):
        dom_finding(
            tmp_db,
            scan_id=scan_id,
            page_id=page_id,
            rule_id="many-elements",
            help_text="Many elements",
            target=f"#el-{n}",
            impact="moderate",
            wcag_sc="1.3.1",
            wcag_level="A",
        )
        _review(tmp_db, scan_id, [_finding_id(tmp_db, scan_id, f"#el-{n}")], notes=f"Element {n}.")
    wb = _render(tmp_db, scan_id)
    narrative = str(wb["Summary"]["A4"].value)

    # Every reported element is still listed, past the template's last row.
    assert len(_issue_rows(wb)) == 50
    assert "reported 50 elements with 1 problem" in narrative
    assert "The counts below include the first 48 rows only." in narrative
    assert "2 more rows are listed" in narrative
    assert "Blocker" in narrative


def test_coversheet_says_only_what_the_scan_knows(tmp_db: sqlite3.Connection) -> None:
    scan_id = seed_rich_scan(tmp_db)
    cover = _render(tmp_db, scan_id)["Coversheet"]

    assert cover["B3"].value == datetime(2026, 4, 22)
    # Tester(s) is the team's to fill, even when the evaluation names a reviewer.
    assert cover["B4"].value == _template()["Coversheet"]["B4"].value
    assert cover["B8"].value == "example.org"
    assert cover["B17"].value == "https://example.org/"
    assert cover["B17"].hyperlink.target == "https://example.org/"
    # A sign-in scan is not recorded as one, so Login info is the team's to fill.
    assert cover["B18"].value == _template()["Coversheet"]["B18"].value
    assert "Rule check (axe)" in cover["B19"].value


def test_coversheet_keeps_the_prompts_a_person_must_answer(tmp_db: sqlite3.Connection) -> None:
    scan_id = _scan_with_real_findings(tmp_db)
    cover = _render(tmp_db, scan_id)["Coversheet"]
    template = _template()["Coversheet"]

    for cell in ("B4", "B9", "B10", "B11", "B12", "B13"):
        assert cover[cell].value == template[cell].value


def test_never_reads_another_scan(tmp_db: sqlite3.Connection) -> None:
    scan_a = _rich(tmp_db)
    scan_b = _rich(tmp_db)
    page = repo.upsert_page(
        tmp_db,
        scan_id=scan_b,
        url_normalized="https://example.org/only-in-b",
        status_code=200,
        title="Only in scan B",
        render_mode="js",
        html_hash="7" * 64,
    )
    dom_finding(
        tmp_db,
        scan_id=scan_b,
        page_id=page,
        rule_id="only-in-b",
        help_text="Only in scan B",
        target="#b",
        impact="critical",
        wcag_sc="1.3.1",
        wcag_level="A",
    )
    _review(tmp_db, scan_b, [_finding_id(tmp_db, scan_b, "#b")], notes="Scan B's own note.")
    wb = _render(tmp_db, scan_a)

    text = _all_text(wb)
    assert "Only in scan B" not in text
    assert "Scan B's own note." not in text
    assert len(_issue_rows(wb)) == 4


def test_renders_when_reviews_cannot_be_read(tmp_db: sqlite3.Connection) -> None:
    """A database without the review table yet still exports, with no rows."""
    scan_id = _rich(tmp_db)
    tmp_db.execute("DROP TABLE occurrence_reviews")
    wb = _render(tmp_db, scan_id)
    assert _issue_rows(wb) == []


def test_cli_writes_the_its_review(
    tmp_db: sqlite3.Connection, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    scan_id = seed_rich_scan(tmp_db)
    tmp_db.commit()
    db_path = Path(str(tmp_db.execute("PRAGMA database_list").fetchone()[2]))
    data_dir = tmp_path / "data"
    settings = Settings(
        data_dir=data_dir, db_path=db_path, blob_dir=data_dir / "blobs", log_dir=data_dir / "logs"
    )
    monkeypatch.setattr(cli, "get_settings", lambda: settings)

    result = CliRunner().invoke(cli.app, ["export", str(scan_id), "--format", "its"])

    assert result.exit_code == 0, result.output
    written = data_dir / "exports" / f"scan_{scan_id}.its.xlsx"
    assert load_workbook(written).sheetnames == _template().sheetnames
