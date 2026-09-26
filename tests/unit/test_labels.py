"""Click-Through is the one name a reader sees for the interaction probe.

Stored data, API fields and code keep ``interaction`` / ``dom_state`` as
stable internal keys, but the method ledger, every export and the CLI help
must say "Click-Through" (exactly that capitalization) and never fall back to
the older "DOM state" / "interaction state" wording.
"""

from __future__ import annotations

import io
import json
import re
import sqlite3
from dataclasses import replace

import pytest
from openpyxl import load_workbook
from support.export_render import EXPORT_FORMATS, render_unlabeled
from support.rich_scan import seed_rich_scan
from typer.testing import CliRunner

from audit import labels
from audit.cli import app
from audit.exports.interaction_coverage import InteractionCoverage
from audit.web.server import _methods_used

# Old user-facing names. Underscored keys (``dom_state_count``) are machine
# fields, not prose, so they deliberately do not match. "Click-Through
# states" was retired for "views opened by clicking", words anyone can follow.
_RETIRED = re.compile(
    r"dom[ -]states?|interaction states?|click[- ]?through states?", re.IGNORECASE
)
# Any spelling of the feature name; each one must be the canonical form.
_ANY_SPELLING = re.compile(r"click[- ]?through", re.IGNORECASE)


def assert_reader_text_is_clean(text: str, *, where: str) -> None:
    retired = _RETIRED.findall(text)
    assert not retired, f"{where} still says {sorted(set(retired))}"
    spellings = set(_ANY_SPELLING.findall(text)) - {labels.CLICK_THROUGH}
    assert not spellings, f"{where} spells the feature as {sorted(spellings)}"


def workbook_text(data: bytes) -> str:
    wb = load_workbook(io.BytesIO(data))
    parts: list[str] = list(wb.sheetnames)
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            parts.extend(str(value) for value in row if value is not None)
    return "\n".join(parts)


def test_labels_are_the_canonical_spelling() -> None:
    assert labels.CLICK_THROUGH == "Click-Through"
    assert labels.CLICK_THROUGH_STATE == "view opened by clicking"
    assert labels.CLICK_THROUGH_STATES == "views opened by clicking"
    assert labels.CLICK_THROUGH_STATES_LABEL == "Views opened by clicking"
    assert labels.click_through_states(1) == "1 view opened by clicking"
    assert labels.click_through_states(3) == "3 views opened by clicking"


@pytest.fixture
def rich_scan(tmp_db: sqlite3.Connection) -> int:
    return seed_rich_scan(tmp_db)


@pytest.mark.parametrize("export_format", EXPORT_FORMATS)
def test_every_export_names_the_feature_click_through(
    tmp_db: sqlite3.Connection, rich_scan: int, export_format: str
) -> None:
    rendered = render_unlabeled(tmp_db, rich_scan, export_format)
    text = workbook_text(rendered) if isinstance(rendered, bytes) else rendered
    assert_reader_text_is_clean(text, where=export_format)
    if export_format in {"audit", "xlsx"}:
        # The two formats that report Click-Through coverage must name it.
        assert labels.CLICK_THROUGH_STATES in text
    if export_format == "xlsx":
        assert labels.CLICK_THROUGH in load_workbook(io.BytesIO(rendered)).sheetnames


def test_method_ledger_row_is_labelled_click_through() -> None:
    methods = _methods_used(
        {
            "status": "completed",
            "page_count": 2,
            "config_json": json.dumps(
                {
                    "interaction_checks_enabled": True,
                    "interaction_coverage_version": 2,
                    "interaction_safety_version": 1,
                }
            ),
            "interaction_pages_probed": 2,
            "interaction_states_total": 1,
        },
        {
            "rendered_pages": 2,
            "analyzed_images": 0,
            "discovered_images": 0,
            "interaction_controls": 4,
            "interaction_operated": 4,
        },
    )
    method = next(item for item in methods if item["key"] == "interaction")
    assert method["label"] == labels.CLICK_THROUGH
    assert "1 view opened by clicking" in method["result"]
    for field in ("label", "result", "description", "caveat"):
        assert_reader_text_is_clean(str(method[field]), where=f"method {field}")


_RAN = InteractionCoverage(
    enabled=True,
    ledger_recorded=True,
    pages_probed=2,
    states_total=5,
    controls_found=16,
    controls_operated=12,
    blocked_controls=1,
    dialogs_stuck=1,
    findings_revealed=2,
)


@pytest.mark.parametrize(
    "coverage",
    [
        replace(_RAN, enabled=False),
        replace(_RAN, ledger_recorded=False),
        replace(_RAN, controls_found=0, controls_operated=0),
        _RAN,
    ],
    ids=["off", "no-ledger", "no-controls", "ran"],
)
def test_interaction_coverage_sentences_say_click_through(coverage: InteractionCoverage) -> None:
    assert coverage.status_line.startswith(labels.CLICK_THROUGH)
    for sentence in (coverage.status_line, *coverage.caveats):
        assert_reader_text_is_clean(sentence, where="interaction coverage")


def test_ran_status_line_counts_click_through_states() -> None:
    assert "reaching 5 views opened by clicking" in _RAN.status_line


def test_cli_help_names_the_skip_interaction_feature() -> None:
    result = CliRunner().invoke(app, ["crawl", "--help"], env={"COLUMNS": "200"})
    assert result.exit_code == 0, result.output
    assert "--skip-interaction" in result.output
    assert f"Skip {labels.CLICK_THROUGH}" in result.output
